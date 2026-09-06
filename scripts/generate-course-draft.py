"""
Builds content-templates/course-content-draft.xlsx from the JSON written by
generate-course-draft.ts. Run the TS script first (it reads the live
catalogue and drafts overview/outcomes/curriculum suggestions); this step
only formats what it wrote — it never invents or edits draft content itself.

    pnpm dlx tsx scripts/generate-course-draft.ts
    python3 scripts/generate-course-draft.py

Every row lands with draftStatus = "NEEDS REVIEW". Nothing in this workbook
is read by the website or by the import workflow (scripts/import-course-content.py
only reads columns matching its known template shape) until someone marks a
row APPROVED and its content is manually carried into
content-templates/course-content-template.xlsx (or a future draft-import
step is explicitly requested).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "content-templates"

with open(OUT_DIR / "course-content-draft.json", encoding="utf-8") as f:
    rows = json.load(f)

# Verified/source columns (unchanged from the master template) + new draft columns.
COLUMNS = [
    "slug", "courseCode", "courseName", "category", "subcategory", "courseType",
    "duration", "curriculum",
    "mode", "level", "overview", "outcomes", "eligibility", "certification", "fees",
    "featured", "status", "notes",
    "overviewDraft", "outcomesDraft", "curriculumDraft", "draftStatus", "draftFlags",
]

DRAFT_COLUMNS = {"overviewDraft", "outcomesDraft", "curriculumDraft", "draftStatus", "draftFlags"}
VERIFIED_TRACKED = {"mode", "level", "overview", "outcomes", "eligibility", "certification", "fees"}


def flatten_curriculum_modules(modules):
    """[{title, topics}] -> 'Title >> topic | topic || Title2 >> topic' (matches the importer's richer-curriculum syntax)."""
    if not modules:
        return ""
    return " || ".join(f"{m['title']} >> " + " | ".join(m["topics"]) for m in modules)


def cell_value(row, key):
    if key == "curriculum":
        return " | ".join(row["curriculum"])
    if key == "curriculumDraft":
        return flatten_curriculum_modules(row.get("curriculumDraft"))
    if key == "outcomes" or key == "outcomesDraft" or key == "eligibility":
        value = row.get(key)
        return " | ".join(value) if value else ""
    if key == "draftFlags":
        return " | ".join(row.get("draftFlags", []))
    if key == "featured":
        return "TRUE" if row["featured"] else "FALSE"
    value = row.get(key)
    if value is None:
        return ""
    return str(value)


wb = Workbook()

# --- Sheet 1: Instructions --------------------------------------------------
instructions = wb.active
instructions.title = "Instructions"
instructions.sheet_view.showGridLines = False
instructions.column_dimensions["A"].width = 100

title_font = Font(name="Arial", size=14, bold=True)
body_font = Font(name="Arial", size=11)
bold_font = Font(name="Arial", size=11, bold=True)

lines = [
    ("SRS Academy — Course Content Draft Sheet (for review)", title_font),
    ("", body_font),
    ("What this is", bold_font),
    ("- Editorial DRAFT suggestions for overview, outcomes and curriculum presentation,", body_font),
    ("  generated only from course name, category, subcategory, duration and the", body_font),
    ("  existing verified curriculum topics. Nothing here has been published — this", body_font),
    ("  workbook does not change the live website and was not written to enrichment.ts.", body_font),
    ("", body_font),
    ("How to use it", bold_font),
    ("- Review each row's overviewDraft / outcomesDraft / curriculumDraft.", body_font),
    ("- Set draftStatus to APPROVED once you're satisfied a row is accurate, or", body_font),
    ("  REJECTED to discard it. Rows stay NEEDS REVIEW until you change them.", body_font),
    ("- Nothing marked NEEDS REVIEW may be imported into the live enrichment source.", body_font),
    ("- draftFlags calls out anything uncertain: too few curriculum topics to draft", body_font),
    ("  outcomes safely, no obvious grouping for curriculumDraft, or a course-type", body_font),
    ("  word (Certificate / Diploma / Advanced Diploma / Training Program) that must", body_font),
    ("  never be read as proof of accreditation, government recognition, university", body_font),
    ("  affiliation, internship placement, or professional licensing.", body_font),
    ("- Columns mode / level / eligibility / certification / fees are intentionally", body_font),
    ("  blank unless already verified — this pass never drafts or infers them.", body_font),
    ("- overview / outcomes stay blank in this sheet when a draft was generated for", body_font),
    ("  them (see overviewDraft / outcomesDraft instead); where overview/outcomes", body_font),
    ("  already show a value, that value is already verified and was left untouched —", body_font),
    ("  no draft was generated for that field.", body_font),
    ("", body_font),
    ("Getting an approved draft into the live site", bold_font),
    ("- Approving a row here does not publish it. Copy the approved overview/outcomes/", body_font),
    ("  curriculum text into content-templates/course-content-template.xlsx (or an", body_font),
    ("  edited copy of it) and run the existing two-step import workflow:", body_font),
    ("    python3 scripts/import-course-content.py <edited.xlsx>          (preview)", body_font),
    ("    python3 scripts/import-course-content.py <edited.xlsx> --apply  (publish)", body_font),
]
for i, (text, font) in enumerate(lines, start=1):
    cell = instructions.cell(row=i, column=1, value=text)
    cell.font = font
    cell.alignment = Alignment(wrap_text=False, vertical="top")

total = len(rows)
with_overview_draft = sum(1 for r in rows if r.get("overviewDraft"))
with_outcomes_draft = sum(1 for r in rows if r.get("outcomesDraft"))
with_curriculum_draft = sum(1 for r in rows if r.get("curriculumDraft"))
flagged = sum(1 for r in rows if r.get("draftFlags"))

row_after = len(lines) + 2
instructions.cell(row=row_after, column=1, value="Generation summary").font = bold_font
instructions.cell(
    row=row_after + 1,
    column=1,
    value=(
        f"{total} courses processed — overviewDraft: {with_overview_draft}, "
        f"outcomesDraft: {with_outcomes_draft}, curriculumDraft: {with_curriculum_draft}, "
        f"rows with a flag: {flagged}"
    ),
).font = body_font

# --- Sheet 2: Courses --------------------------------------------------------
ws = wb.create_sheet("Courses")
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"

header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
draft_header_fill = PatternFill(start_color="1E5A8E", end_color="1E5A8E", fill_type="solid")
verified_tracked_fill = PatternFill(start_color="92722A", end_color="92722A", fill_type="solid")
draft_cell_fill = PatternFill(start_color="E8F1FB", end_color="E8F1FB", fill_type="solid")
flag_fill = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
cell_font = Font(name="Arial", size=10)

for col_idx, name in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=col_idx, value=name)
    c.font = header_font
    if name in DRAFT_COLUMNS:
        c.fill = draft_header_fill
    elif name in VERIFIED_TRACKED:
        c.fill = verified_tracked_fill
    else:
        c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

draft_col_indexes = {i for i, name in enumerate(COLUMNS, start=1) if name in DRAFT_COLUMNS}
flags_col_index = COLUMNS.index("draftFlags") + 1
status_col_index = COLUMNS.index("draftStatus") + 1
status_col_letter = get_column_letter(status_col_index)

for r, row in enumerate(rows, start=2):
    for c_idx, key in enumerate(COLUMNS, start=1):
        value = cell_value(row, key)
        cell = ws.cell(row=r, column=c_idx, value=value if value != "" else None)
        cell.font = cell_font
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c_idx == flags_col_index and value:
            cell.fill = flag_fill
        elif c_idx in draft_col_indexes and value:
            cell.fill = draft_cell_fill

# Restrict draftStatus to the three allowed values via a dropdown.
dv = DataValidation(type="list", formula1='"NEEDS REVIEW,APPROVED,REJECTED"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"{status_col_letter}2:{status_col_letter}{len(rows) + 1}")

widths = {
    "slug": 34, "courseCode": 12, "courseName": 34, "category": 22, "subcategory": 22,
    "courseType": 14, "duration": 12, "curriculum": 50,
    "mode": 14, "level": 14, "overview": 34, "outcomes": 34, "eligibility": 34,
    "certification": 30, "fees": 18, "featured": 10, "status": 10, "notes": 26,
    "overviewDraft": 46, "outcomesDraft": 46, "curriculumDraft": 55,
    "draftStatus": 14, "draftFlags": 46,
}
for i, name in enumerate(COLUMNS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 18)
ws.row_dimensions[1].height = 32

wb.save(OUT_DIR / "course-content-draft.xlsx")
print(f"Wrote {OUT_DIR / 'course-content-draft.xlsx'} ({len(rows)} rows)")
